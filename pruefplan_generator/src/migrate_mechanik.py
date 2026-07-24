"""Importiert die primäre Regel-Datenbank in die SQLite-Datenbank:
- 'Mechanik' und 'Produktspezifikationen' aus Preisliste_Mechanik_v2.xlsm (aktuelle
  Preise; ersetzt die frühere, meist leere Preisspalten enthaltende Quelle)
- 'PA_Database' und 'ADMIN' weiterhin aus HL_KV_Monitoring_PA_ab_Sept(MAK).xlsm

Idempotent: löscht und befüllt die Zieltabellen bei jedem Lauf neu.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from src.db import clear_table, get_connection, init_schema, normalize_warengruppe

SOURCE_DIR = Path(__file__).resolve().parent.parent.parent
MECHANIK_SOURCE_FILE = SOURCE_DIR / "Preisliste_Mechanik_v2 (1).xlsm"
MONITORING_SOURCE_FILE = SOURCE_DIR / "HL_KV_Monitoring_PA_ab_Sept(MAK).xlsm"
# Rückwärtskompatibler Alias (frühere Ausbaustufe)
SOURCE_FILE = MONITORING_SOURCE_FILE


def _to_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def migrate_mechanik(conn, wb) -> int:
    """Liest das Blatt 'Mechanik' aus Preisliste_Mechanik_v2.xlsm (23 Spalten,
    u.a. mehrere Kosten-Varianten -- wir übernehmen 'Kosten VP 1.Produkt' (aktuell,
    nicht 'alt') und 'Kosten TP')."""
    ws = wb["Mechanik"]
    clear_table(conn, "mechanik_regeln")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    count = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        (
            lidl_bereich,
            _hilfsspalte1,
            _hilfsspalte2,
            lidl_warengruppe,
            warengruppe,
            produkt,
            bemerkungen,
            _vergangene_projekte,
            trivial,
            norm_ek_ppm,
            anzahl_muster,
            kez_anforderungen,
            _kosten_vp_1_alt,
            kosten_vp,
            _kosten_vp_2,
            kosten_tp,
            ffu,
            kosten_ffu,
            stiwa,
            kosten_stiwa,
            _istkosten_mechanik,
            _laborminuten_mechanik,
            _istkosten_labor_ffu,
        ) = (list(row) + [None] * 23)[:23]

        if not produkt:
            continue

        warengruppe_code = normalize_warengruppe(warengruppe) or normalize_warengruppe(lidl_warengruppe)

        conn.execute(
            """
            INSERT INTO mechanik_regeln (
                lidl_bereich, lidl_warengruppe, warengruppe_code, produkt,
                bemerkungen, trivial, norm_ek_ppm, anzahl_muster, kez_anforderungen,
                kosten_vp, kosten_tp, ffu, kosten_ffu, stiwa, kosten_stiwa, bewertung_uv
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                _to_str(lidl_bereich),
                _to_str(lidl_warengruppe),
                warengruppe_code,
                _to_str(produkt),
                _to_str(bemerkungen),
                _to_str(trivial),
                _to_str(norm_ek_ppm),
                _to_str(anzahl_muster),
                _to_str(kez_anforderungen),
                _to_float(kosten_vp),
                _to_float(kosten_tp),
                _to_str(ffu),
                _to_float(kosten_ffu),
                _to_str(stiwa),
                _to_float(kosten_stiwa),
                "",
            ),
        )
        count += 1
    conn.commit()
    return count


def migrate_produktspezifikationen(conn, wb) -> int:
    """Liest das Blatt 'Produktspezifikationen' (Materialanforderungskatalog)."""
    ws = wb["Produktspezifikationen"]
    clear_table(conn, "produktspezifikationen")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    count = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        (
            parameter,
            produkt,
            mak_paket,
            norm,
            material_code,
            laufzeit,
            _formel_laufzeit,
            kosten,
            bewertung,
            anzahl_proben,
            _formel_bewertung,
            _kosten_bewertung,
            _formel_staffelung,
            _kosten_staffelung,
            gesamtkosten,
            vk_preis,
            pruefstelle,
            kontakt,
            bemerkung,
            materialverkaufstext,
        ) = (list(row) + [None] * 20)[:20]

        if not parameter:
            continue

        conn.execute(
            """
            INSERT INTO produktspezifikationen (
                parameter, produkt, mak_paket, norm, material_code, laufzeit,
                kosten, bewertung, anzahl_proben, gesamtkosten, vk_preis,
                pruefstelle, kontakt, bemerkung, materialverkaufstext
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                _to_str(parameter),
                _to_str(produkt),
                _to_str(mak_paket),
                _to_str(norm),
                _to_str(material_code),
                _to_str(laufzeit),
                _to_float(kosten),
                _to_str(bewertung),
                _to_str(anzahl_proben),
                _to_float(gesamtkosten),
                _to_float(vk_preis),
                _to_str(pruefstelle),
                _to_str(kontakt),
                _to_str(bemerkung),
                _to_str(materialverkaufstext),
            ),
        )
        count += 1
    conn.commit()
    return count


def migrate_pa_referenzdaten(conn, wb) -> int:
    ws = wb["PA_Database"]
    clear_table(conn, "pa_referenzdaten")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    count = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        (
            ian_charge,
            artikelbezeichnung,
            artikelkategorie,
            ian_vorgaenger,
            warengruppe,
            lieferant,
            pruefumfang,
            qualitaet,
            material,
            markenreferenz,
            referenzpruefung,
        ) = (list(row) + [None] * 11)[:11]

        if not ian_charge:
            continue

        conn.execute(
            """
            INSERT INTO pa_referenzdaten (
                ian_charge, artikelbezeichnung, artikelkategorie, ian_vorgaenger,
                warengruppe, lieferant, pruefumfang, qualitaet, material,
                markenreferenz, referenzpruefung
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                _to_str(ian_charge),
                _to_str(artikelbezeichnung),
                _to_str(artikelkategorie),
                _to_str(ian_vorgaenger),
                _to_str(warengruppe),
                _to_str(lieferant),
                _to_str(pruefumfang),
                _to_str(qualitaet),
                _to_str(material),
                _to_str(markenreferenz),
                _to_str(referenzpruefung),
            ),
        )
        count += 1
    conn.commit()
    return count


def migrate_warengruppe_mapping(conn, wb) -> int:
    ws = wb["ADMIN"]
    clear_table(conn, "warengruppe_mapping")
    rows = list(ws.iter_rows(min_row=2, max_col=2, values_only=True))
    count = 0
    seen = set()
    for row in rows:
        if row is None:
            continue
        neu, alt = (list(row) + [None, None])[:2]
        if not neu:
            continue
        code = normalize_warengruppe(neu)
        if not code or code in seen:
            continue
        seen.add(code)
        conn.execute(
            "INSERT INTO warengruppe_mapping (warengruppe_code, alte_warengruppe) VALUES (?, ?)",
            (code, _to_str(alt)),
        )
        count += 1
    conn.commit()
    return count


def run(
    mechanik_source_file: Path = MECHANIK_SOURCE_FILE,
    monitoring_source_file: Path = MONITORING_SOURCE_FILE,
) -> dict:
    conn = get_connection()
    init_schema(conn)

    mechanik_wb = openpyxl.load_workbook(mechanik_source_file, data_only=True, read_only=True, keep_vba=True)
    monitoring_wb = openpyxl.load_workbook(monitoring_source_file, data_only=True, read_only=True, keep_vba=True)

    result = {
        "mechanik_regeln": migrate_mechanik(conn, mechanik_wb),
        "produktspezifikationen": migrate_produktspezifikationen(conn, mechanik_wb),
        "pa_referenzdaten": migrate_pa_referenzdaten(conn, monitoring_wb),
        "warengruppe_mapping": migrate_warengruppe_mapping(conn, monitoring_wb),
    }
    conn.close()
    return result


if __name__ == "__main__":
    stats = run()
    for table, n in stats.items():
        print(f"{table}: {n} Zeilen importiert")
