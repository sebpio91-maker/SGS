"""Importiert die sekundäre (Fallback-) Regel-Datenbank aus Normenauswahl.xlsm
(Blätter 'Datenbank', 'Sonderposten', 'NORMEN') in die SQLite-Datenbank.

Idempotent: löscht und befüllt die Zieltabellen bei jedem Lauf neu.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from src.db import clear_table, get_connection, init_schema

SOURCE_DIR = Path(__file__).resolve().parent.parent.parent
SOURCE_FILE = SOURCE_DIR / "Normenauswahl.xlsm"


def _to_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def migrate_normen_regeln(conn, wb) -> int:
    ws = wb["Datenbank"]
    clear_table(conn, "normen_regeln")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    count = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        (
            kategorie,
            _spalte1,
            produktart,
            _spalte2,
            zielgruppe,
            _spalte3,
            einsatzort,
            _spalte4,
            bereich,
            _spalte5,
            produkt,
            code,
            lidl_warengruppe,
            normen,
            material,
            verweise,
            preis,
            laborzeit_min,
            laborkosten,
            bemerkungen,
        ) = (list(row) + [None] * 20)[:20]

        if not produkt:
            continue

        conn.execute(
            """
            INSERT OR IGNORE INTO normen_regeln (
                kategorie, produktart, zielgruppe, einsatzort, bereich, produkt,
                code, lidl_warengruppe, normen, material, verweise,
                preis, laborzeit_min, laborkosten, bemerkungen
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                _to_str(kategorie),
                _to_str(produktart),
                _to_str(zielgruppe),
                _to_str(einsatzort),
                _to_str(bereich),
                _to_str(produkt),
                _to_str(code),
                _to_str(lidl_warengruppe),
                _to_str(normen),
                _to_str(material),
                _to_str(verweise),
                _to_float(preis),
                _to_float(laborzeit_min),
                _to_float(laborkosten),
                _to_str(bemerkungen),
            ),
        )
        count += 1
    conn.commit()
    return count


def migrate_sonderposten(conn, wb) -> int:
    ws = wb["Sonderposten"]
    clear_table(conn, "sonderposten")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    count = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        (kategorie, produktart, zielgruppe, einsatzort, eigenschaft, zusatzkosten, kommentar) = (
            list(row) + [None] * 7
        )[:7]

        if not eigenschaft:
            continue

        conn.execute(
            """
            INSERT INTO sonderposten (
                kategorie, produktart, zielgruppe, einsatzort, eigenschaft,
                zusatzkosten, kommentar
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                _to_str(kategorie),
                _to_str(produktart),
                _to_str(zielgruppe),
                _to_str(einsatzort),
                _to_str(eigenschaft),
                _to_float(zusatzkosten),
                _to_str(kommentar),
            ),
        )
        count += 1
    conn.commit()
    return count


def migrate_normen_texte(conn, wb) -> int:
    """Das Blatt 'NORMEN' enthält pro Norm einen Block:
    Zeile 1: Normname (z.B. 'DIN EN 581-1')
    Zeile 3: Spaltenköpfe ('Kapitel', 'Überschrift', 'Prüfungsrelevant', 'Inhalt')
    Zeile 4+: Kapitel-Zeilen, bis eine neue Norm-Kopfzeile (nur Spalte A gefüllt,
              Spalte B/C/D leer und Wert nicht 'Kapitel') den nächsten Block einleitet.
    """
    ws = wb["NORMEN"]
    clear_table(conn, "normen_texte")
    rows = list(ws.iter_rows(values_only=True))

    count = 0
    current_norm = ""
    i = 0
    while i < len(rows):
        row = rows[i]
        if row is None or all(v is None for v in row):
            i += 1
            continue
        a, b, c, d = (list(row) + [None] * 4)[:4]
        if a == "Kapitel" and b == "Überschrift":
            i += 1
            continue
        if b is None and c is None and d is None and a:
            current_norm = _to_str(a)
            i += 1
            continue

        kapitel = _to_str(a)
        ueberschrift = _to_str(b)
        pruefungsrelevant = _to_str(c)
        inhalt = _to_str(d)
        if kapitel or ueberschrift:
            conn.execute(
                """
                INSERT INTO normen_texte (
                    norm_name, kapitel, ueberschrift, pruefungsrelevant, inhalt
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (current_norm, kapitel, ueberschrift, pruefungsrelevant, inhalt),
            )
            count += 1
        i += 1
    conn.commit()
    return count


def run(source_file: Path = SOURCE_FILE) -> dict:
    conn = get_connection()
    init_schema(conn)
    wb = openpyxl.load_workbook(source_file, data_only=True, read_only=True, keep_vba=True)
    result = {
        "normen_regeln": migrate_normen_regeln(conn, wb),
        "sonderposten": migrate_sonderposten(conn, wb),
        "normen_texte": migrate_normen_texte(conn, wb),
    }
    conn.close()
    return result


if __name__ == "__main__":
    stats = run()
    for table, n in stats.items():
        print(f"{table}: {n} Zeilen importiert")
