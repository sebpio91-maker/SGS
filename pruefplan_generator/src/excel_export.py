"""Excel-Export des Prüfplans im Layout von HL_KV_Monitoring...xlsm!Auswahl_LIDL:
Kopfdaten-Block + Positionstabelle (Setbestandteil/Produkt/Prüfung/Anzahl/Kosten) + Summe.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from src.models import PruefplanErgebnis

_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)


def build_workbook(ergebnis: PruefplanErgebnis) -> Workbook:
    artikel = ergebnis.artikel
    wb = Workbook()
    ws = wb.active
    ws.title = "Prüfplan"

    ws["A1"] = "Prüfplan"
    ws["A1"].font = _TITLE_FONT

    kopf_felder = [
        ("IAN_Charge", artikel.ian_charge),
        ("Artikelbezeichnung", artikel.artikelbezeichnung),
        ("Bereich", artikel.warenbereich),
        ("Warengruppe", artikel.warengruppe),
        ("Artikelkategorie", artikel.artikelkategorie),
        ("Lieferant", artikel.lieferant),
        ("Herkunftsland", artikel.herkunftsland),
        ("Prüfumfang (Prüfauftrag)", ", ".join(artikel.pruefumfang)),
        ("Quelle", artikel.quelle_datei),
    ]
    row = 3
    for label, value in kopf_felder:
        ws.cell(row=row, column=1, value=label).font = _HEADER_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    tabellen_start = row
    headers = ["Setbestandteil", "Produkt", "Prüfung", "Anzahl", "Kosten [EUR]", "Bemerkung", "Quelle"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=tabellen_start, column=col, value=header)
        cell.font = _HEADER_FONT

    row = tabellen_start + 1
    for pos in ergebnis.positionen:
        ws.cell(row=row, column=1, value=pos.setbestandteil)
        ws.cell(row=row, column=2, value=pos.produkt)
        ws.cell(row=row, column=3, value=pos.pruefung)
        ws.cell(row=row, column=4, value=pos.anzahl)
        ws.cell(row=row, column=5, value=pos.kosten)
        ws.cell(row=row, column=6, value=pos.bemerkung)
        ws.cell(row=row, column=7, value=pos.quelle)
        row += 1

    ws.cell(row=row, column=1, value="Gesamtkosten").font = _HEADER_FONT
    ws.cell(row=row, column=5, value=ergebnis.gesamtkosten).font = _HEADER_FONT

    for col, width in zip("ABCDEFG", [22, 30, 42, 10, 14, 30, 14]):
        ws.column_dimensions[col].width = width
    ws["B3"].alignment = Alignment(wrap_text=True)

    return wb


def export_bytes(ergebnis: PruefplanErgebnis) -> bytes:
    wb = build_workbook(ergebnis)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_file(ergebnis: PruefplanErgebnis, path: str) -> None:
    wb = build_workbook(ergebnis)
    wb.save(path)
