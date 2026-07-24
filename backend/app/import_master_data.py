"""Importiert die realen Stammdaten aus den drei Referenz-Excel-Dateien.

    python -m app.import_master_data \
        --normenauswahl /pfad/zu/Normenauswahl.xlsm \
        --produktspezifikationen /pfad/zu/Produktspezifikationen.xlsx \
        --mechanik /pfad/zu/KV_Monitoring.xlsm

Jede Quelle ist optional einzeln angebbar. Die drei Zieltabellen
`norm_special_items`, `product_spec_requirements` und
`lidl_warengruppe_rules` werden bei jedem Import vollständig ersetzt
(Referenzdaten, kein Nutzer-generierter Inhalt). `norm_lookup_rules` wird
per eindeutigem `code` upserted.
"""

import argparse

import openpyxl

from app.db.session import SessionLocal
from app.models.master_data import (
    LidlWarengruppeRule,
    NormLookupRule,
    NormSpecialItem,
    ProductSpecRequirement,
)


def _num(value):
    if isinstance(value, (int, float)):
        return value
    return None


def _str(value):
    if value is None:
        return None
    return str(value).strip() or None


def _flag(value) -> bool:
    return isinstance(value, str) and value.strip().lower() == "x"


def import_normenauswahl(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    db = SessionLocal()
    try:
        ws = wb["Datenbank"]
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            kategorie = row[0]
            if not kategorie:
                continue
            (
                kategorie,
                _s1,
                produktart,
                _s2,
                zielgruppe,
                _s3,
                einsatzort,
                _s4,
                bereich,
                _s5,
                produkt,
                code,
                lidl_warengruppe,
                normen,
                material,
                verweise,
                preis,
                zeit_labor_min,
                kosten_labor,
                bemerkungen,
            ) = row

            if not normen or not code:
                continue

            rule = db.query(NormLookupRule).filter_by(code=code).one_or_none()
            if rule is None:
                rule = NormLookupRule(code=code)
                db.add(rule)

            rule.kategorie = kategorie
            rule.produktart = produktart
            rule.zielgruppe = _str(zielgruppe)
            rule.einsatzort = _str(einsatzort)
            rule.bereich = _str(bereich)
            rule.produkt = produkt
            rule.lidl_warengruppe = _str(lidl_warengruppe)
            rule.normen = normen
            rule.material = _str(material)
            rule.verweise = _str(verweise)
            rule.preis = _num(preis)
            rule.zeit_labor_min = _num(zeit_labor_min)
            rule.kosten_labor = _num(kosten_labor)
            rule.bemerkungen = _str(bemerkungen)
            imported += 1

        db.commit()
        print(f"NormLookupRule: {imported} Zeilen importiert/aktualisiert.")

        ws = wb["Sonderposten"]
        db.query(NormSpecialItem).delete()
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            kategorie, produktart, zielgruppe, einsatzort, eigenschaft, zusatzkosten, kommentar = row
            if not eigenschaft:
                continue
            db.add(
                NormSpecialItem(
                    kategorie=_str(kategorie),
                    produktart=_str(produktart),
                    zielgruppe=_str(zielgruppe),
                    einsatzort=_str(einsatzort),
                    eigenschaft=eigenschaft,
                    zusatzkosten=_num(zusatzkosten),
                    kommentar=_str(kommentar),
                )
            )
            imported += 1
        db.commit()
        print(f"NormSpecialItem: {imported} Zeilen importiert (Tabelle neu befüllt).")
    finally:
        db.close()


def import_produktspezifikationen(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    db = SessionLocal()
    try:
        ws = wb["Produktspezifikationen"]
        db.query(ProductSpecRequirement).delete()
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            (
                parameter,
                produkt,
                mak,
                norm,
                material,
                laufzeit,
                formel_laufzeit,
                kosten,
                bewertung,
                anzahl_proben,
                formel_bewertung,
                kosten_bewertung,
                formel_staffelung,
                kosten_staffelung,
                gesamtkosten,
                vk_preis,
                pruefstelle,
                kontakt,
                bemerkung,
                materialverkaufstext,
            ) = row[:20]

            if not parameter:
                continue

            db.add(
                ProductSpecRequirement(
                    parameter=parameter,
                    produkt=_str(produkt),
                    mak=_str(mak),
                    norm=_str(norm),
                    material_code=_str(material),
                    laufzeit=_num(laufzeit),
                    formel_laufzeit=_str(formel_laufzeit),
                    kosten=_num(kosten),
                    bewertung=_str(bewertung),
                    anzahl_proben=_str(anzahl_proben),
                    formel_bewertung=_str(formel_bewertung),
                    kosten_bewertung=_num(kosten_bewertung),
                    formel_staffelung=_str(formel_staffelung),
                    kosten_staffelung=_num(kosten_staffelung),
                    gesamtkosten=_str(gesamtkosten),
                    vk_preis=_num(vk_preis),
                    pruefstelle=_str(pruefstelle),
                    kontakt=_str(kontakt),
                    bemerkung=_str(bemerkung),
                    materialverkaufstext=_str(materialverkaufstext),
                )
            )
            imported += 1
        db.commit()
        print(f"ProductSpecRequirement: {imported} Zeilen importiert (Tabelle neu befüllt).")
    finally:
        db.close()


def import_mechanik(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    db = SessionLocal()
    try:
        ws = wb["Mechanik"]
        db.query(LidlWarengruppeRule).delete()
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            (
                lidl_bereich,
                lidl_warengruppe,
                warengruppe_code,
                _hilfsspalte,
                produkt,
                bemerkungen,
                trivial,
                norm_ek_ppm,
                anzahl_muster,
                kez_anforderung,
                kosten_vp,
                ffu_code,
                kosten_ffu,
                stiwa_referenz,
                kosten_stiwa,
                bewertung_uv,
            ) = row[:16]

            if not lidl_bereich:
                continue

            db.add(
                LidlWarengruppeRule(
                    lidl_bereich=lidl_bereich,
                    lidl_warengruppe=_str(lidl_warengruppe),
                    warengruppe_code=_num(warengruppe_code),
                    produkt=_str(produkt),
                    bemerkungen=_str(bemerkungen),
                    trivial=_flag(trivial),
                    norm_ek_ppm=_str(norm_ek_ppm),
                    anzahl_muster=_str(anzahl_muster),
                    kez_anforderung=_flag(kez_anforderung),
                    kosten_vp=_num(kosten_vp),
                    ffu_code=_str(ffu_code),
                    kosten_ffu=_num(kosten_ffu),
                    stiwa_referenz=_str(stiwa_referenz),
                    kosten_stiwa=_num(kosten_stiwa),
                    bewertung_uv=_str(bewertung_uv),
                )
            )
            imported += 1
        db.commit()
        print(f"LidlWarengruppeRule: {imported} Zeilen importiert (Tabelle neu befüllt).")
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--normenauswahl", help="Pfad zu Normenauswahl.xlsm")
    parser.add_argument("--produktspezifikationen", help="Pfad zu Produktspezifikationen.xlsx")
    parser.add_argument("--mechanik", help="Pfad zu KV_Monitoring.xlsm (Blatt 'Mechanik')")
    args = parser.parse_args()

    if args.normenauswahl:
        import_normenauswahl(args.normenauswahl)
    if args.produktspezifikationen:
        import_produktspezifikationen(args.produktspezifikationen)
    if args.mechanik:
        import_mechanik(args.mechanik)

    if not any([args.normenauswahl, args.produktspezifikationen, args.mechanik]):
        parser.print_help()
